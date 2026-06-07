"""One-page portfolio snapshot generator (UC-18).

Produces a PDF (via ``reportlab``) or DOCX (via ``python-docx``)
formatted to the Linear design system spirit — heavy typographic
hierarchy on a clean canvas — but executed in print-friendly inks since
the export is intended for grant applications and conference packets.
"""
from __future__ import annotations

import io
import logging
from datetime import date
from typing import Iterable, List, Sequence
from uuid import UUID

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.core.exceptions import NotFoundError, ValidationFailure
from app.db.models import (
    AcademicBackground,
    Publication,
    User,
    UserExpertiseTag,
)
from app.schemas import ExportSnapshotRequest

logger = logging.getLogger(__name__)


class ExportService:
    """Assembles and renders the customized one-page snapshot."""

    def __init__(self, session: AsyncSession) -> None:
        self._session = session

    async def generate(
        self, *, user_id: UUID, request: ExportSnapshotRequest
    ) -> tuple[bytes, str, str]:
        user = await self._load_user(user_id)

        tags = self._filter_tags(user, request)
        publications = self._filter_publications(user, request)
        background = self._filter_background(user, request)

        if request.use_defaults and not (tags or publications or background):
            # UC-18 alt flow: nothing validated yet — use the top items.
            tags = self._default_tags(user)
            publications = self._default_publications(user)
            background = self._default_background(user)

        if not (tags or publications or background):
            # UC-18 exception — empty profile selection.
            raise ValidationFailure(
                "Snapshot must include at least one expertise tag, publication, "
                "or background record.",
            )

        if request.format == "pdf":
            buffer = self._render_pdf(user, tags, publications, background)
            mime = "application/pdf"
            filename = f"{self._slug(user.full_name)}-snapshot.pdf"
        else:
            buffer = self._render_docx(user, tags, publications, background)
            mime = (
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            )
            filename = f"{self._slug(user.full_name)}-snapshot.docx"
        return buffer, mime, filename

    # ------------------------------------------------------------------ Loaders
    async def _load_user(self, user_id: UUID) -> User:
        stmt = (
            select(User)
            .where(User.id == user_id)
            .options(
                selectinload(User.expertise_links).selectinload(
                    UserExpertiseTag.tag
                ),
                selectinload(User.publications),
                selectinload(User.academic_background),
            )
        )
        user = (await self._session.execute(stmt)).scalar_one_or_none()
        if user is None:
            raise NotFoundError("User not found.")
        return user

    # ------------------------------------------------------------------ Filters
    @staticmethod
    def _filter_tags(user: User, request: ExportSnapshotRequest) -> List[UserExpertiseTag]:
        if not request.include_tag_ids:
            return []
        wanted = set(request.include_tag_ids)
        return [link for link in user.expertise_links if link.tag.id in wanted]

    @staticmethod
    def _filter_publications(
        user: User, request: ExportSnapshotRequest
    ) -> List[Publication]:
        if not request.include_publication_ids:
            return []
        wanted = set(request.include_publication_ids)
        return [p for p in user.publications if p.id in wanted]

    @staticmethod
    def _filter_background(
        user: User, request: ExportSnapshotRequest
    ) -> List[AcademicBackground]:
        if not request.include_background_ids:
            return []
        wanted = set(request.include_background_ids)
        return [b for b in user.academic_background if b.id in wanted]

    @staticmethod
    def _default_tags(user: User) -> List[UserExpertiseTag]:
        validated = [link for link in user.expertise_links if link.validated]
        chosen = validated or list(user.expertise_links)
        return sorted(chosen, key=lambda l: l.confidence, reverse=True)[:8]

    @staticmethod
    def _default_publications(user: User) -> List[Publication]:
        return sorted(
            user.publications,
            key=lambda p: (p.publication_year or 0),
            reverse=True,
        )[:5]

    @staticmethod
    def _default_background(user: User) -> List[AcademicBackground]:
        return sorted(
            user.academic_background,
            key=lambda b: b.start_date,
            reverse=True,
        )[:5]

    # ------------------------------------------------------------------ PDF
    @staticmethod
    def _render_pdf(
        user: User,
        tags: Sequence[UserExpertiseTag],
        publications: Sequence[Publication],
        background: Sequence[AcademicBackground],
    ) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph,
            SimpleDocTemplate,
            Spacer,
        )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=18 * mm,
            bottomMargin=18 * mm,
        )
        styles = getSampleStyleSheet()
        h1 = ParagraphStyle(
            "ExpertiseTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            spaceAfter=2,
        )
        meta = ParagraphStyle(
            "ExpertiseMeta",
            parent=styles["Normal"],
            textColor="#555555",
            fontSize=10,
            spaceAfter=12,
        )
        section = ParagraphStyle(
            "ExpertiseSection",
            parent=styles["Heading2"],
            fontSize=12,
            spaceBefore=10,
            spaceAfter=4,
        )
        body = ParagraphStyle(
            "ExpertiseBody",
            parent=styles["Normal"],
            fontSize=10,
            leading=14,
        )

        flow = [
            Paragraph(user.full_name, h1),
            Paragraph(
                f"{user.department or 'Faculty Member'} &middot; {user.email}",
                meta,
            ),
        ]

        if tags:
            flow.append(Paragraph("Expertise Tags", section))
            label_line = " &nbsp;|&nbsp; ".join(
                f"<b>{link.tag.canonical_label}</b>" for link in tags
            )
            flow.append(Paragraph(label_line, body))

        if publications:
            flow.append(Paragraph("Publications", section))
            for pub in publications:
                line = (
                    f"&bull; <b>{pub.title or '(Untitled)'}</b> "
                    f"<i>{pub.venue or ''}</i> {pub.publication_year or ''}".strip()
                )
                flow.append(Paragraph(line, body))

        if background:
            flow.append(Paragraph("Academic Background", section))
            for record in background:
                end = record.end_date.isoformat() if record.end_date else "present"
                line = (
                    f"&bull; <b>{record.title}</b> &mdash; "
                    f"{record.organization or ''} "
                    f"({record.start_date.isoformat()} &ndash; {end})"
                )
                flow.append(Paragraph(line, body))
                if record.description:
                    flow.append(Paragraph(record.description, body))

        flow.append(Spacer(1, 6))
        doc.build(flow)
        return buffer.getvalue()

    # ------------------------------------------------------------------ DOCX
    @staticmethod
    def _render_docx(
        user: User,
        tags: Sequence[UserExpertiseTag],
        publications: Sequence[Publication],
        background: Sequence[AcademicBackground],
    ) -> bytes:
        from docx import Document
        from docx.shared import Pt

        document = Document()
        heading = document.add_heading(user.full_name, level=0)
        for run in heading.runs:
            run.font.size = Pt(20)
        document.add_paragraph(
            f"{user.department or 'Faculty Member'} · {user.email}"
        )

        if tags:
            document.add_heading("Expertise Tags", level=2)
            document.add_paragraph(
                ", ".join(link.tag.canonical_label for link in tags)
            )

        if publications:
            document.add_heading("Publications", level=2)
            for pub in publications:
                line = pub.title or "(Untitled)"
                if pub.venue:
                    line += f" — {pub.venue}"
                if pub.publication_year:
                    line += f" ({pub.publication_year})"
                document.add_paragraph(line, style="List Bullet")

        if background:
            document.add_heading("Academic Background", level=2)
            for record in background:
                end = record.end_date.isoformat() if record.end_date else "present"
                document.add_paragraph(
                    f"{record.title} — {record.organization or ''} "
                    f"({record.start_date.isoformat()} – {end})",
                    style="List Bullet",
                )
                if record.description:
                    document.add_paragraph(record.description)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _slug(value: str) -> str:
        return "-".join(part for part in value.lower().split() if part) or "snapshot"
