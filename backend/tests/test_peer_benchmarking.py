"""Focused PEER-01 evidence, ordering, and representative-file regressions."""
from __future__ import annotations

import asyncio
import io
from datetime import datetime, timezone
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from app.core.exceptions import ValidationFailure
from app.api.routes_benchmarking import _to_run_out
from app.db.models import BenchmarkType
from app.services.benchmarking import BenchmarkingService
from app.services.document_extract import extract_text_from_upload
from app.services.document_representation import SourceSection


def test_peer_api_contract_exposes_descriptive_source_evidence() -> None:
    evidence = SimpleNamespace(
        domain_label="synthetic.pdf — Distributed Systems",
        displacement_score=None,
        recommendation=None,
        peer_filename="synthetic.pdf",
        detected_section="Distributed Systems",
        source_excerpt="Consensus, replication, and fault tolerance.",
        nearest_staff_id=None,
        nearest_staff_name=None,
        nearest_cluster_label=None,
        nearest_cluster_distance=None,
        coverage_classification=None,
        review_status="source_overview",
    )
    run = SimpleNamespace(
        id=uuid4(),
        benchmark_type=BenchmarkType.PEER,
        created_at=datetime(2026, 10, 20, tzinfo=timezone.utc),
        narrative=None,
        source_payload={"mode": "descriptive_curriculum_overview"},
        visualization_payload=None,
        white_spaces=[evidence],
    )

    response = _to_run_out(run)

    assert response.white_spaces[0].peer_filename == "synthetic.pdf"
    assert response.white_spaces[0].displacement_score is None
    assert response.white_spaces[0].nearest_staff_id is None
    assert response.white_spaces[0].coverage_classification is None
    assert response.white_spaces[0].review_status == "source_overview"


def test_peer_overview_does_not_call_embeddings_staff_matching_or_llm(monkeypatch) -> None:
    service = BenchmarkingService(SimpleNamespace())
    captured: dict[str, object] = {}

    async def no_op() -> None:
        return None

    async def unexpected(*args, **kwargs):
        raise AssertionError("Peer overview must not run analytical comparison")

    async def persist(**kwargs):
        captured.update(kwargs)
        return SimpleNamespace(id=uuid4())

    service._acquire_run_lock = no_op  # type: ignore[method-assign]
    service._load_internal_expertise = unexpected  # type: ignore[method-assign]
    service._persist_run = persist  # type: ignore[method-assign]
    monkeypatch.setattr("app.services.benchmarking.embed_texts", unexpected)
    monkeypatch.setattr("app.services.benchmarking.synthesize_narrative", unexpected)
    section = SourceSection(
        id="section-1",
        filename="synthetic.pdf",
        section="Distributed Systems",
        excerpt="Consensus, replication, and fault tolerance.",
        included_by_default=True,
        token_count=6,
    )

    asyncio.run(service.run_peer_overview(triggered_by=uuid4(), peer_sections=[section]))

    assert captured["narrative"] is None
    assert captured["visualization"] is None
    evidence = captured["white_spaces"][0]
    assert evidence.displacement_score is None
    assert evidence.nearest_staff_name is None


def test_generated_text_pdf_is_readable() -> None:
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(72, 760, "Year 1 Core Curriculum")
    document.drawString(72, 740, "Students study algorithms, databases, and secure software design.")
    document.save()

    text = asyncio.run(extract_text_from_upload("synthetic-programme.pdf", buffer.getvalue(), "application/pdf"))

    assert "Year 1 Core Curriculum" in text
    assert "secure software design" in text


def test_generated_docx_table_is_preserved_as_text() -> None:
    document = Document()
    document.add_heading("Programme Structure", level=1)
    table = document.add_table(rows=2, cols=3)
    table.rows[0].cells[0].text = "Year"
    table.rows[0].cells[1].text = "Code"
    table.rows[0].cells[2].text = "Course"
    table.rows[1].cells[0].text = "2"
    table.rows[1].cells[1].text = "SYN204"
    table.rows[1].cells[2].text = "Distributed Systems and Security"
    buffer = io.BytesIO()
    document.save(buffer)

    text = asyncio.run(extract_text_from_upload(
        "synthetic-table.docx",
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ))

    assert "Year | Code | Course" in text
    assert "2 | SYN204 | Distributed Systems and Security" in text


def test_generated_extraction_poor_pdf_is_rejected_with_ocr_guidance() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buffer = io.BytesIO()
    writer.write(buffer)

    with pytest.raises(ValidationFailure, match="scanned or image-only PDFs require OCR"):
        asyncio.run(extract_text_from_upload("synthetic-scan.pdf", buffer.getvalue(), "application/pdf"))
