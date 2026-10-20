"""Focused UC-18 administrator staff CV export regressions."""
from __future__ import annotations

import asyncio
from datetime import date
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from docx import Document
from pypdf import PdfReader

from app.api import routes_export
from app.core.exceptions import ForbiddenError, NotFoundError
from app.db.models import UserRole
from app.schemas import ExportSnapshotRequest
from app.services.export import ExportService


def _default_request() -> ExportSnapshotRequest:
    return ExportSnapshotRequest(format="pdf", use_defaults=True)


def test_academic_staff_cannot_export_another_staff_cv() -> None:
    actor = SimpleNamespace(role=UserRole.ACADEMIC_STAFF)

    with pytest.raises(ForbiddenError, match="Only administrators"):
        asyncio.run(
            routes_export.export_staff_snapshot(
                user_id=uuid4(),
                payload=_default_request(),
                session=object(),  # type: ignore[arg-type]
                actor=actor,  # type: ignore[arg-type]
            )
        )


def test_administrator_cannot_export_ineligible_target() -> None:
    class FakeSession:
        async def scalar(self, statement: object) -> None:
            return None

    actor = SimpleNamespace(role=UserRole.FACULTY_ADMINISTRATOR)

    with pytest.raises(NotFoundError, match="profile not found"):
        asyncio.run(
            routes_export.export_staff_snapshot(
                user_id=uuid4(),
                payload=_default_request(),
                session=FakeSession(),  # type: ignore[arg-type]
                actor=actor,  # type: ignore[arg-type]
            )
        )


def test_administrator_exports_target_with_default_selection(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    target_id = uuid4()
    captured: dict[str, object] = {}

    class FakeSession:
        async def scalar(self, statement: object) -> object:
            return target_id

    class FakeExportService:
        def __init__(self, session: object) -> None:
            pass

        async def generate(
            self, *, user_id: object, request: ExportSnapshotRequest
        ) -> tuple[bytes, str, str]:
            captured["user_id"] = user_id
            captured["use_defaults"] = request.use_defaults
            return b"pdf", "application/pdf", "staff-snapshot.pdf"

    monkeypatch.setattr(routes_export, "ExportService", FakeExportService)
    actor = SimpleNamespace(role=UserRole.FACULTY_ADMINISTRATOR)
    response = asyncio.run(
        routes_export.export_staff_snapshot(
            user_id=target_id,
            payload=_default_request(),
            session=FakeSession(),  # type: ignore[arg-type]
            actor=actor,  # type: ignore[arg-type]
        )
    )

    assert response.body == b"pdf"
    assert captured == {"user_id": target_id, "use_defaults": True}


def test_redesigned_cv_templates_render_structured_one_page_content() -> None:
    user = SimpleNamespace(
        full_name="Dr. Amina Rahman",
        department="Department of Software Engineering",
        email="amina.rahman@example.edu",
        phone_number="+60 3-7967 0000",
    )
    tags = [
        SimpleNamespace(tag=SimpleNamespace(canonical_label=label))
        for label in (
            "Responsible Artificial Intelligence",
            "Software Quality Engineering",
            "Digital Health Systems",
            "Human-Computer Interaction",
            "Sustainable Computing",
            "Data Governance",
        )
    ]
    publications = [
        SimpleNamespace(
            publication_year=2026 - index,
            title=f"Evidence-informed software strategy study {index + 1}",
            venue="International Journal of Computing Strategy",
        )
        for index in range(4)
    ]
    background = [
        SimpleNamespace(
            start_date=date(2018 - index * 3, 1, 1),
            end_date=None if index == 0 else date(2020 - index, 12, 31),
            title=(
                "Associate Professor"
                if index == 0
                else f"Academic appointment {index + 1}"
            ),
            organization="Example University",
            description="Research leadership, programme development, and collaboration.",
        )
        for index in range(3)
    ]

    pdf_bytes = ExportService._render_pdf(user, tags, publications, background)
    docx_bytes = ExportService._render_docx(user, tags, publications, background)

    pdf = PdfReader(BytesIO(pdf_bytes))
    document = Document(BytesIO(docx_bytes))
    document_text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    ) + "\n" + "\n".join(
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    )

    pdf_text = pdf.pages[0].extract_text() or ""
    pdf_lines = pdf_text.splitlines()
    assert len(pdf.pages) == 1
    assert "Core Expertise" in pdf_text
    assert "DEPARTMENT    Department of Software Engineering" in pdf_lines
    assert "EMAIL    amina.rahman@example.edu" in pdf_lines
    assert "PHONE    +60 3-7967 0000" in pdf_lines
    assert "CORE EXPERTISE" in document_text
    assert "Dr. Amina Rahman" in document_text
    assert len(document.tables[0].rows) == 3
    assert [row.cells[1].text for row in document.tables[0].rows] == [
        "DEPARTMENT  Department of Software Engineering",
        "EMAIL  amina.rahman@example.edu",
        "PHONE  +60 3-7967 0000",
    ]
